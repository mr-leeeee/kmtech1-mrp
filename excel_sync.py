import openpyxl
import os
import shutil
import logging
import msvcrt
from datetime import datetime
from typing import Optional
from mrp_engine import (MRPSystem, Item, FinishedGood, BOMItem, ProductionLog,
                        Transaction, PurchaseOrder, FGTransaction)

logger = logging.getLogger(__name__)

def find_data_start(ws, key_column: int = 1, header_keyword: str = None, max_scan: int = 50) -> int:
    """시트에서 데이터 시작 행을 동적으로 탐색 (헤더 다음 행부터)."""
    for r in range(1, max_scan + 1):
        val = str(ws.cell(row=r, column=key_column).value or "").strip()
        if header_keyword:
            if header_keyword in val:
                return r + 1
        elif val and not val.startswith(('─', '═', '━', '═', '---', '===')):
            if r > 1:
                prev_val = str(ws.cell(row=r - 1, column=key_column).value or "").strip()
                if not prev_val or prev_val.startswith(('─', '═', '━', '═', '---', '===')):
                    return r
    return 6

def find_last_row(ws, key_column: int = 1, start_row: int = 6, max_scan: int = 200,
                   section_keywords: list = None) -> int:
    """데이터가 있는 마지막 행을 탐색. section_keywords가 있으면 해당 키워드出现 시 중단."""
    last_row = start_row
    for r in range(start_row, start_row + max_scan):
        val = str(ws.cell(row=r, column=key_column).value or "").strip()
        if section_keywords and val and any(kw in val for kw in section_keywords):
            break
        if ws.cell(row=r, column=key_column).value is not None and val:
            last_row = r
        elif r > last_row + 5:
            break
    return last_row

def load_from_excel(filepath: str) -> MRPSystem:
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Excel file not found: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    mrp = MRPSystem()

    if "품목_기준정보" in wb.sheetnames:
        ws = wb["품목_기준정보"]
        start = find_data_start(ws, key_column=1, header_keyword="코드")
        if start <= 1:
            start = 6
        end = find_last_row(ws, key_column=1, start_row=start,
                           section_keywords=["완제품"])
        
        for r in range(start, end + 1):
            code = str(ws.cell(row=r, column=1).value or "").strip()
            name = str(ws.cell(row=r, column=2).value or "").strip()
            category = str(ws.cell(row=r, column=3).value or "").strip()
            unit = str(ws.cell(row=r, column=4).value or "EA").strip()
            try:
                base_stock = float(ws.cell(row=r, column=5).value or 0)
            except (ValueError, TypeError):
                base_stock = 0
            try:
                safety_stock = float(ws.cell(row=r, column=6).value or 0)
            except (ValueError, TypeError):
                safety_stock = 0
            try:
                target_stock = float(ws.cell(row=r, column=7).value or (safety_stock * 2))
            except (ValueError, TypeError):
                target_stock = safety_stock * 2
            try:
                lead_time = int(ws.cell(row=r, column=8).value or 5)
            except (ValueError, TypeError):
                lead_time = 5
            note = str(ws.cell(row=r, column=9).value or "").strip()
            if code and name:
                mrp.add_item(Item(
                    code=code, name=name, category=category, unit=unit,
                    base_stock=base_stock, safety_stock=safety_stock,
                    target_stock=target_stock, lead_time=lead_time, note=note
                ))

        fg_ws = wb["품목_기준정보"]
        fg_start = find_data_start(fg_ws, key_column=1, header_keyword="완제품코드")
        if fg_start <= 1:
            fg_start = 24
        fg_end = find_last_row(fg_ws, key_column=1, start_row=fg_start)
        
        for r in range(fg_start, fg_end + 1):
            fg_code = str(fg_ws.cell(row=r, column=1).value or "").strip()
            fg_name = str(fg_ws.cell(row=r, column=2).value or "").strip()
            spec = str(fg_ws.cell(row=r, column=3).value or "").strip()
            unit = str(fg_ws.cell(row=r, column=4).value or "Set").strip()
            note = str(fg_ws.cell(row=r, column=5).value or "").strip()
            if fg_code and fg_code.startswith("FG") and fg_name:
                mrp.add_finished_good(FinishedGood(
                    code=fg_code, name=fg_name, spec=spec, unit=unit, note=note
                ))

    if "BOM_정규형" in wb.sheetnames:
        ws = wb["BOM_정규형"]
        start = find_data_start(ws, key_column=1, header_keyword="완제품")
        if start <= 1:
            start = 5
        end = find_last_row(ws, key_column=1, start_row=start)
        
        for r in range(start, end + 1):
            fg_code = str(ws.cell(row=r, column=1).value or "").strip()
            fg_name = str(ws.cell(row=r, column=2).value or "").strip()
            part_cat = str(ws.cell(row=r, column=3).value or "").strip()
            part_name = str(ws.cell(row=r, column=4).value or "").strip()
            qty_val = ws.cell(row=r, column=5).value
            unit = str(ws.cell(row=r, column=6).value or "EA").strip()
            note = str(ws.cell(row=r, column=7).value or "").strip()
            try:
                qty_float = float(qty_val) if qty_val is not None else 0
            except (ValueError, TypeError):
                continue
            if fg_code and part_name and qty_float > 0:
                mrp.add_bom_item(BOMItem(
                    fg_code=fg_code, fg_name=fg_name, part_category=part_cat,
                    part_name=part_name, unit_qty=qty_float, unit=unit, note=note
                ))

    if "생산_실적일지" in wb.sheetnames:
        ws = wb["생산_실적일지"]
        start = find_data_start(ws, key_column=1, header_keyword="일자")
        if start <= 1:
            start = 4
        end = find_last_row(ws, key_column=4, start_row=start)
        
        for r in range(start, end + 1):
            date_val = ws.cell(row=r, column=1).value
            slip = str(ws.cell(row=r, column=2).value or "").strip()
            line = str(ws.cell(row=r, column=3).value or "").strip()
            fg_code = str(ws.cell(row=r, column=4).value or "").strip()
            fg_name = str(ws.cell(row=r, column=5).value or "").strip()
            spec = str(ws.cell(row=r, column=6).value or "").strip()
            p_qty = ws.cell(row=r, column=7).value
            g_qty = ws.cell(row=r, column=8).value
            d_qty = ws.cell(row=r, column=9).value
            worker = str(ws.cell(row=r, column=10).value or "").strip()
            note = str(ws.cell(row=r, column=12).value or "").strip()
            try:
                p_qty_f = float(p_qty) if p_qty is not None else 0
            except (ValueError, TypeError):
                continue
            if fg_code and p_qty_f > 0:
                try:
                    g_qty_f = float(g_qty) if g_qty is not None else p_qty_f
                except (ValueError, TypeError):
                    g_qty_f = p_qty_f
                try:
                    d_qty_f = float(d_qty) if d_qty is not None else 0
                except (ValueError, TypeError):
                    d_qty_f = 0
                d_str = date_val.strftime("%Y-%m-%d") if hasattr(date_val, "strftime") else str(date_val)[:10]
                mrp.production_logs.append(ProductionLog(
                    id=len(mrp.production_logs) + 1, date=d_str, slip_no=slip, line=line,
                    fg_code=fg_code, fg_name=fg_name, spec=spec,
                    prod_qty=p_qty_f, good_qty=g_qty_f,
                    defect_qty=d_qty_f, worker=worker, note=note
                ))
        mrp._next_prod_id = len(mrp.production_logs) + 1

    if "자재_입출고원장" in wb.sheetnames:
        ws = wb["자재_입출고원장"]
        start = find_data_start(ws, key_column=1, header_keyword="일자")
        if start <= 1:
            start = 4
        end = find_last_row(ws, key_column=5, start_row=start)
        
        for r in range(start, end + 1):
            date_val = ws.cell(row=r, column=1).value
            slip = str(ws.cell(row=r, column=2).value or "").strip()
            tx_type = str(ws.cell(row=r, column=3).value or "").strip()
            part_cat = str(ws.cell(row=r, column=4).value or "").strip()
            part_name = str(ws.cell(row=r, column=5).value or "").strip()
            unit = str(ws.cell(row=r, column=6).value or "EA").strip()
            try:
                in_q = float(ws.cell(row=r, column=7).value or 0)
            except (ValueError, TypeError):
                in_q = 0
            try:
                out_q = float(ws.cell(row=r, column=8).value or 0)
            except (ValueError, TypeError):
                out_q = 0
            cust = str(ws.cell(row=r, column=9).value or "").strip()
            note = str(ws.cell(row=r, column=10).value or "").strip()
            worker = str(ws.cell(row=r, column=11).value or "").strip()
            if part_name and (in_q > 0 or out_q > 0):
                d_str = date_val.strftime("%Y-%m-%d") if hasattr(date_val, "strftime") else str(date_val)[:10]
                mrp.transactions.append(Transaction(
                    id=len(mrp.transactions) + 1, date=d_str, slip_no=slip, tx_type=tx_type,
                    part_category=part_cat, part_name=part_name, unit=unit,
                    in_qty=in_q, out_qty=out_q, customer=cust, note=note, worker=worker
                ))
        mrp._next_tx_id = len(mrp.transactions) + 1

    if "발주_미입고관리" in wb.sheetnames:
        ws = wb["발주_미입고관리"]
        start = find_data_start(ws, key_column=1, header_keyword="발주")
        if start <= 1:
            start = 4
        end = find_last_row(ws, key_column=4, start_row=start)
        
        for r in range(start, end + 1):
            date_val = ws.cell(row=r, column=1).value
            slip = str(ws.cell(row=r, column=2).value or "").strip()
            p_code = str(ws.cell(row=r, column=3).value or "").strip()
            p_name = str(ws.cell(row=r, column=4).value or "").strip()
            unit = str(ws.cell(row=r, column=5).value or "EA").strip()
            ord_q = ws.cell(row=r, column=6).value
            rec_q = ws.cell(row=r, column=7).value
            exp_val = ws.cell(row=r, column=9).value
            cust = str(ws.cell(row=r, column=10).value or "").strip()
            stat = str(ws.cell(row=r, column=11).value or "").strip()
            note = str(ws.cell(row=r, column=12).value or "").strip()
            try:
                ord_qty = float(ord_q) if ord_q is not None else 0
            except (ValueError, TypeError):
                continue
            if p_name and ord_qty > 0:
                try:
                    recv_qty = float(rec_q) if rec_q is not None else 0
                except (ValueError, TypeError):
                    recv_qty = 0
                d_str = date_val.strftime("%Y-%m-%d") if hasattr(date_val, "strftime") else str(date_val)[:10]
                e_str = exp_val.strftime("%Y-%m-%d") if hasattr(exp_val, "strftime") else str(exp_val)[:10]
                pend_qty = max(ord_qty - recv_qty, 0.0)
                mrp.purchase_orders.append(PurchaseOrder(
                    id=len(mrp.purchase_orders) + 1, date=d_str, po_no=slip, part_code=p_code,
                    part_name=p_name, unit=unit, order_qty=ord_qty, recv_qty=recv_qty,
                    pending_qty=pend_qty, exp_date=e_str, customer=cust, status=stat, note=note
                ))
        mrp._next_po_id = len(mrp.purchase_orders) + 1

    if "완제품_출하관리" in wb.sheetnames:
        ws = wb["완제품_출하관리"]
        start = find_data_start(ws, key_column=1, header_keyword="일자")
        if start <= 1:
            start = 4
        end = find_last_row(ws, key_column=3, start_row=start)

        for r in range(start, end + 1):
            date_val = ws.cell(row=r, column=1).value
            slip = str(ws.cell(row=r, column=2).value or "").strip()
            fg_code = str(ws.cell(row=r, column=3).value or "").strip()
            fg_name = str(ws.cell(row=r, column=4).value or "").strip()
            unit = str(ws.cell(row=r, column=5).value or "Set").strip()
            out_val = ws.cell(row=r, column=6).value
            cust = str(ws.cell(row=r, column=7).value or "").strip()
            note = str(ws.cell(row=r, column=8).value or "").strip()
            worker = str(ws.cell(row=r, column=9).value or "").strip()
            try:
                out_qty = float(out_val) if out_val is not None else 0
            except (ValueError, TypeError):
                continue
            if fg_code and out_qty > 0:
                d_str = date_val.strftime("%Y-%m-%d") if hasattr(date_val, "strftime") else str(date_val)[:10]
                mrp.fg_transactions.append(FGTransaction(
                    id=len(mrp.fg_transactions) + 1, date=d_str, slip_no=slip,
                    fg_code=fg_code, fg_name=fg_name, unit=unit, out_qty=out_qty,
                    customer=cust, note=note, worker=worker
                ))
        mrp._next_fg_tx_id = len(mrp.fg_transactions) + 1

    if "재고현황_대시보드" in wb.sheetnames:
        ws = wb["재고현황_대시보드"]
        start = find_data_start(ws, key_column=1, header_keyword="완제품")
        if start <= 1:
            start = 40
        end = min(start + 10, 50)
        
        for r in range(start, end + 1):
            fg_code = str(ws.cell(row=r, column=1).value or "").strip()
            plan_val = ws.cell(row=r, column=3).value
            if not fg_code or not fg_code.startswith("FG"):
                continue
            try:
                plan_float = float(plan_val)
            except (ValueError, TypeError):
                continue
            if plan_val is not None:
                mrp.production_plans[fg_code] = plan_float

    return mrp


class ExcelLock:
    def __init__(self, filepath: str):
        self.lock_path = filepath + ".lock"
        self.lock_file = None

    def __enter__(self):
        for attempt in range(3):
            try:
                self.lock_file = open(self.lock_path, 'w')
                msvcrt.locking(self.lock_file.fileno(), msvcrt.LK_NBLCK, 1)
                return self
            except IOError:
                if self.lock_file:
                    self.lock_file.close()
                    self.lock_file = None
                if attempt < 2:
                    logger.warning("Excel 파일 잠김 (시도 %d/3). 2초 대기...", attempt + 1)
                    import time
                    time.sleep(2)
                else:
                    raise RuntimeError("Excel 파일이 다른 프로그램(Excel 등)에서 열려 있습니다. 닫은 후 다시 시도해주세요.")

    def __exit__(self, exc_type, exc_val, exc_tb):
        try:
            msvcrt.locking(self.lock_file.fileno(), msvcrt.LK_UNLCK, 1)
        except:
            pass
        self.lock_file.close()
        try:
            os.remove(self.lock_path)
        except:
            pass


def create_backup(filepath: str) -> str:
    if not os.path.exists(filepath):
        return filepath
    
    backup_root = os.path.join(os.path.dirname(filepath), "backup")
    os.makedirs(backup_root, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    date_str = datetime.now().strftime("%Y%m%d")
    daily_dir = os.path.join(backup_root, date_str)
    os.makedirs(daily_dir, exist_ok=True)
    
    filename = os.path.basename(filepath)
    backup_path = os.path.join(daily_dir, f"{filename}.{timestamp}.bak")
    shutil.copy2(filepath, backup_path)
    logger.info("백업 생성: %s", backup_path)
    
    now = datetime.now()
    for entry in os.listdir(backup_root):
        entry_path = os.path.join(backup_root, entry)
        if os.path.isdir(entry_path):
            try:
                entry_date = datetime.strptime(entry, "%Y%m%d")
                if (now - entry_date).days > 7:
                    shutil.rmtree(entry_path)
                    logger.info("오래된 백업 폴더 삭제: %s", entry_path)
            except ValueError:
                continue
    
    return backup_path


def save_to_excel(mrp: MRPSystem, filepath: str):
    with ExcelLock(filepath):
        create_backup(filepath)
        wb = openpyxl.load_workbook(filepath, data_only=False)

        if "품목_기준정보" in wb.sheetnames:
            ws = wb["품목_기준정보"]
            start = find_data_start(ws, key_column=1, header_keyword="코드")
            if start <= 1:
                start = 6
            end = find_last_row(ws, key_column=1, start_row=start,
                               section_keywords=["완제품"])
            
            for r in range(start, end + 1):
                code = str(ws.cell(row=r, column=1).value or "").strip()
                item = mrp.items_by_code.get(code)
                if item:
                    ws.cell(row=r, column=5).value = item.base_stock
                    ws.cell(row=r, column=6).value = item.safety_stock
                    ws.cell(row=r, column=7).value = item.target_stock
                    ws.cell(row=r, column=8).value = item.lead_time

        if "생산_실적일지" in wb.sheetnames:
            ws = wb["생산_실적일지"]
            for r in range(4, 200):
                has_data = any(ws.cell(row=r, column=c).value for c in [1, 2, 4, 7])
                if not has_data and r > 10:
                    break
                if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith(('--', '---', '==')):
                    continue
                for c in [1, 2, 3, 4, 7, 8, 9, 10, 12]:
                    ws.cell(row=r, column=c).value = None

            for idx, log in enumerate(mrp.production_logs):
                r = 4 + idx
                ws.cell(row=r, column=1).value = log.date
                ws.cell(row=r, column=2).value = log.slip_no
                ws.cell(row=r, column=3).value = log.line
                ws.cell(row=r, column=4).value = log.fg_code
                ws.cell(row=r, column=7).value = log.prod_qty
                ws.cell(row=r, column=8).value = log.good_qty
                ws.cell(row=r, column=9).value = log.defect_qty
                ws.cell(row=r, column=10).value = log.worker
                ws.cell(row=r, column=12).value = log.note

        if "자재_입출고원장" in wb.sheetnames:
            ws = wb["자재_입출고원장"]
            for r in range(4, 200):
                has_data = any(ws.cell(row=r, column=c).value for c in [1, 2, 5, 7, 8])
                if not has_data and r > 10:
                    break
                if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith(('--', '---', '==')):
                    continue
                for c in [1, 2, 3, 5, 7, 8, 9, 10, 11]:
                    ws.cell(row=r, column=c).value = None

            for idx, tx in enumerate(mrp.transactions):
                r = 4 + idx
                ws.cell(row=r, column=1).value = tx.date
                ws.cell(row=r, column=2).value = tx.slip_no
                ws.cell(row=r, column=3).value = tx.tx_type
                ws.cell(row=r, column=5).value = tx.part_name
                ws.cell(row=r, column=7).value = tx.in_qty
                ws.cell(row=r, column=8).value = tx.out_qty
                ws.cell(row=r, column=9).value = tx.customer
                ws.cell(row=r, column=10).value = tx.note
                ws.cell(row=r, column=11).value = tx.worker

        if "발주_미입고관리" in wb.sheetnames:
            ws = wb["발주_미입고관리"]
            for r in range(4, 200):
                has_data = any(ws.cell(row=r, column=c).value for c in [1, 2, 4, 6, 7])
                if not has_data and r > 10:
                    break
                if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith(('--', '---', '==')):
                    continue
                for c in [1, 2, 4, 6, 7, 9, 10, 12]:
                    ws.cell(row=r, column=c).value = None

            for idx, po in enumerate(mrp.purchase_orders):
                r = 4 + idx
                ws.cell(row=r, column=1).value = po.date
                ws.cell(row=r, column=2).value = po.po_no
                ws.cell(row=r, column=4).value = po.part_name
                ws.cell(row=r, column=6).value = po.order_qty
                ws.cell(row=r, column=7).value = po.recv_qty
                ws.cell(row=r, column=9).value = po.exp_date
                ws.cell(row=r, column=10).value = po.customer
                ws.cell(row=r, column=12).value = po.note

        fg_out_sheetname = "완제품_출하관리"
        if fg_out_sheetname in wb.sheetnames:
            ws = wb[fg_out_sheetname]
            for r in range(4, 200):
                has_data = any(ws.cell(row=r, column=c).value for c in [1, 2, 3, 6])
                if not has_data and r > 10:
                    break
                if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith(('--', '---', '==')):
                    continue
                for c in [1, 2, 3, 4, 5, 6, 7, 8, 9]:
                    ws.cell(row=r, column=c).value = None
        else:
            ws = wb.create_sheet(fg_out_sheetname)
            ws.cell(row=1, column=1).value = "완제품 출하 관리"
            ws.cell(row=4, column=1).value = "일자"
            ws.cell(row=4, column=2).value = "전표번호"
            ws.cell(row=4, column=3).value = "완제품코드"
            ws.cell(row=4, column=4).value = "완제품명"
            ws.cell(row=4, column=5).value = "단위"
            ws.cell(row=4, column=6).value = "출하수량"
            ws.cell(row=4, column=7).value = "거래처"
            ws.cell(row=4, column=8).value = "비고"
            ws.cell(row=4, column=9).value = "담당자"

        for idx, fg_tx in enumerate(mrp.fg_transactions):
            r = 4 + idx
            ws.cell(row=r, column=1).value = fg_tx.date
            ws.cell(row=r, column=2).value = fg_tx.slip_no
            ws.cell(row=r, column=3).value = fg_tx.fg_code
            ws.cell(row=r, column=4).value = fg_tx.fg_name
            ws.cell(row=r, column=5).value = fg_tx.unit
            ws.cell(row=r, column=6).value = fg_tx.out_qty
            ws.cell(row=r, column=7).value = fg_tx.customer
            ws.cell(row=r, column=8).value = fg_tx.note
            ws.cell(row=r, column=9).value = fg_tx.worker

        if "재고현황_대시보드" in wb.sheetnames:
            ws = wb["재고현황_대시보드"]
            start = find_data_start(ws, key_column=1, header_keyword="완제품")
            if start <= 1:
                start = 40
            end = min(start + 10, 50)
            
            for r in range(start, end + 1):
                fg_code = str(ws.cell(row=r, column=1).value or "").strip()
                if fg_code in mrp.production_plans:
                    ws.cell(row=r, column=3).value = mrp.production_plans[fg_code]

        wb.save(filepath)
        logger.info("엑셀 저장 완료: %s", filepath)